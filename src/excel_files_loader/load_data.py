import csv
import pandas as pd
from openpyxl.reader.excel import load_workbook

file_location = "/data/YachtCharterData.xlsx"

workbook = load_workbook(file_location)

print(workbook.sheetnames)
worksheet = workbook['Yacht Charter Data']
print(worksheet.max_row)
print(worksheet.max_column)
header_row = worksheet[1]
column_names = [cell.value for cell in header_row]
# print(column_names)

exported_csv_location = "yachts.csv"
# with open(exported_csv_location, "w", newline="", encoding="utf-8") as file:
#     csv_writer = csv.writer(file)
#     for row in worksheet.iter_rows(values_only=True):
#         csv_writer.writerow(row)

df = pd.read_csv(exported_csv_location)
print(df.head())
print(df.shape)
print(df.dtypes)
print(df[df["visit_id"]=="V110"])

# Endpoint na podstawie tego pliku


